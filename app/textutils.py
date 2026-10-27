"""Pure helpers shared by the retriever, catalog and router (from 5.7, no I/O state)."""
from __future__ import annotations

import csv
import functools
import hashlib
import io
import json
import os
import re
import unicodedata
from pathlib import Path
from typing import Any

from app.config import EXCEL_EXTS, OTHER_SECTION, SECTION_LABEL, settings

_CODE_RE = re.compile(r"\b[A-Za-z]{1,6}-?\d{1,5}[A-Za-z0-9\-]*\b|\b\d{5,}\b")
STEM_STOPWORDS = {"drawing", "dibujo", "exploded", "view", "vista", "explosionada",
                  "explosionado", "despiece", "plano", "page", "pagina", "image", "scan"}


def simple_tokenize(text: str) -> list[str]:
    return re.findall(r"[^\W_]+", (text or "").lower())


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def extract_codes(text: str) -> set[str]:
    """Normalized model/article codes ('VDV-25' -> 'vdv25')."""
    return {c for c in (norm(m) for m in _CODE_RE.findall(text or "")) if len(c) >= 3}


def raw_codes(text: str) -> dict[str, str]:
    """normalized code -> first spelling seen (e.g. 'vdv25' -> 'VDV-25')."""
    out: dict[str, str] = {}
    for m in _CODE_RE.findall(text or ""):
        n = norm(m)
        if len(n) >= 3 and n not in out:
            out[n] = m
    return out


def fold(s: str) -> str:
    """lower-case, German umlauts -> ae/oe/ue, accents removed, punctuation -> spaces."""
    s = (s or "").lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue") \
        .replace("ß", "ss")
    s = "".join(ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s).split())


def mkey(s: str) -> str:
    """Loose key for names: 'Schütze', 'Schuetze' and 'schutze' all give 'schutze'."""
    return fold(s).replace("ue", "u").replace(" ", "")


# =========================================================================================
# Production 1.0 — BM25 analyzer (Spanish + English)
# =========================================================================================
# Accent-free (they go through fold()) Spanish and English stop words. Kept short on
# purpose: words like "no", "sin", "con", "max" or "min" are NOT stop words here because
# they change the meaning of a technical question.
_STOP_ES = frozenset("""
a al algo algun alguna algunas alguno algunos ante antes aquel aquella aquellas aquellos
aqui asi aun cada como cual cuales cualquier cuando de del desde donde dos el ella ellas
ello ellos en entre era eran es esa esas ese eso esos esta estaba estado estan estar este
esto estos fue fueron ha habia han hasta hay la las le les lo los mas me mi mis mucho
muy nos o os otra otras otro otros para pero poco por porque que quien quienes se sea
ser si sido sobre su sus tambien tan tanto te tiene tienen todo todos tu tus un una unas
uno unos usted ustedes y ya yo favor quiero quisiera necesito puedes podria dime dame
""".split())
_STOP_EN = frozenset("""
a about above after again all am an and any are as at be been before being below between
both but by can could did do does doing down during each few for from further had has
have having he her here hers him his how i if in into is it its itself just me more most
my of off on once only or other our out over own same she should so some such than that
the their them then there these they this those through to too under until up very was
we were what when where which while who whom why will with would you your please tell
""".split())
STOPWORDS = _STOP_ES | _STOP_EN

try:                                                        # pip install snowballstemmer
    import snowballstemmer as _snowball                     # type: ignore
except ImportError:                                         # pragma: no cover
    _snowball = None

import threading as _threading

_stem_local = _threading.local()                            # Snowball stemmers keep state


def _stemmers():
    st = getattr(_stem_local, "st", None)
    if st is None:
        st = (_snowball.stemmer("spanish"), _snowball.stemmer("english"))
        _stem_local.st = st
    return st


@functools.lru_cache(maxsize=300_000)
def _stems(word: str) -> tuple[str, ...]:
    """Spanish and English stem of one folded word (both kept when they differ, so a
    Spanish query matches an English manual and vice versa). Cached: the vocabulary is
    small, the corpus is big."""
    if _snowball is None or not settings.bm25_stemming:
        return (word,)
    es, en = (s.stemWord(word) for s in _stemmers())
    return (es,) if es == en else (es, en)


def bm25_tokenize(text: str) -> list[str]:
    """Tokens for the BM25 index AND the BM25 query (both sides must use this).

    'Presión máxima de la válvula VDV-25' ->
        ['presion', 'maxim', 'maxima', 'valvul', 'valvula', 'vdv', '25', 'vdv25']
    * fold(): lower case, accents removed, ä/ö/ü -> ae/oe/ue, punctuation -> space;
    * Spanish / English stop words dropped; single letters dropped (digits kept);
    * words with digits are never stemmed (codes, sizes, pressures);
    * every part code found by extract_codes() is added once in normalised form, so
      'VDV-25', 'VDV 25' and 'vdv25' all match the same chunk."""
    out: list[str] = []
    for w in fold(text).split():
        if w in STOPWORDS or (len(w) == 1 and not w.isdigit()):
            continue
        if any(ch.isdigit() for ch in w):
            out.append(w)
        else:
            out.extend(_stems(w))
    out.extend(sorted(extract_codes(text)))
    return out


def query_numbers(text: str) -> frozenset[str]:
    """Numbers written in a question ('20:1', '1/2', '3000') — two questions that differ
    only in a number are different questions (used by the answer cache)."""
    return frozenset(re.findall(r"\d+(?:[.,:/]\d+)*", text or ""))


def parse_json(raw: str) -> dict:
    raw = (raw or "").strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[-1].rsplit("```", 1)[0]
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            return json.loads(m.group(0))
        raise


def fmt_secs(s: float) -> str:
    return f"{s:.1f}s" if s < 60 else f"{int(s // 60)}m{int(s % 60):02d}s"


def file_hash(path: str | Path) -> str:
    """Content fingerprint (sha1, streamed): detects 'same file, new date'."""
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def caption_from_text(text: str) -> str:
    return text.split("Description:\n", 1)[1].strip() if "Description:\n" in text else ""


def read_csv_table(path: str | Path) -> tuple[list[str], list[list[str]], list[str]]:
    """Table file -> (headers, rows, preamble). Name kept for compatibility: it reads CSV
    AND Excel (.xlsx / .xlsm / .xls).

    CSV: detects encoding, delimiter and skips title rows above the real header.
    Excel: every visible sheet with data is read (cell values as Excel shows them, merged
    cells filled with their value). With several sheets the rows are stacked into ONE
    table with a first column "Hoja" (sheet name) and the union of the sheets' columns, so
    row numbers stay stable for citations and product lookups."""
    if Path(path).suffix.lower() in EXCEL_EXTS:
        sheets = [(name, rows) for name, rows in read_excel_sheets(path) if rows]
        if not sheets:
            return [], [], []
        if len(sheets) == 1:
            return _table_from_rows(sheets[0][1])
        headers: list[str] = ["Hoja"]
        data: list[list[str]] = []
        preamble: list[str] = []
        parts = []
        for name, rows in sheets:
            h, d, pre = _table_from_rows(rows)
            if not h:
                continue
            parts.append((name, h, d))
            preamble += [f"{name}: {x}" for x in pre]
            for col in h:
                if col not in headers:
                    headers.append(col)
        pos = {h: i for i, h in enumerate(headers)}
        for name, h, d in parts:
            for r in d:
                row = [""] * len(headers)
                row[0] = name
                for col, v in zip(h, r):
                    row[pos[col]] = v
                data.append(row)
        return headers, data, preamble

    raw = Path(path).read_bytes()
    text = ""
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:20000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        delim = ";" if first.count(";") > first.count(",") else ","
    return _table_from_rows(list(csv.reader(io.StringIO(text), delimiter=delim)))


def _table_from_rows(raw_rows: list[list[str]]) -> tuple[list[str], list[list[str]], list[str]]:
    """Grid of cell strings -> (headers, rows, preamble): finds the real header row below
    any title rows, names empty / repeated headers, pads rows, optional forward fill."""
    rows = [[" ".join(str(c).split()) for c in r] for r in raw_rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return [], [], []

    def filled(r: list[str]) -> int:
        return sum(1 for c in r if c)

    max_filled = max(filled(r) for r in rows[:50])
    hdr_i = next((i for i, r in enumerate(rows[:15]) if filled(r) >= max(2, 0.5 * max_filled)), 0)
    preamble = [" ".join(c for c in r if c) for r in rows[:hdr_i]]
    width = max(len(r) for r in rows)
    header_row = rows[hdr_i] + [""] * (width - len(rows[hdr_i]))
    headers, used = [], set()
    for j, h in enumerate(header_row):
        h = h or f"Column {j + 1}"
        base, n = h, 2
        while h in used:
            h, n = f"{base} ({n})", n + 1
        used.add(h)
        headers.append(h)

    data = [r + [""] * (width - len(r)) for r in rows[hdr_i + 1:]]
    for c in range(min(settings.csv_forward_fill_cols, width)):
        last = ""
        for r in data:
            if r[c]:
                last = r[c]
            elif last:
                r[c] = last
    return headers, data, preamble


# ---------------------------------------------------------------------------- Excel
_EXCEL_FULL_MODE_BYTES = 15 * 1024 * 1024    # below this: merged cells are filled


def _excel_value(v: Any, number_format: str = "") -> str:
    """One Excel cell -> the text a person sees (no '12.0', no '0.30000000000000004')."""
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, _dt.datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, _dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)):
        pct = "%" in (number_format or "")
        x = v * 100 if pct else v
        if isinstance(x, float):
            if x != x or x in (float("inf"), float("-inf")):
                return ""
            x = round(x, 10)
            txt = str(int(x)) if x.is_integer() else repr(x)
        else:
            txt = str(x)
        return txt + ("%" if pct else "")
    return str(v).strip()


def read_excel_sheets(path: str | Path) -> list[tuple[str, list[list[str]]]]:
    """[(sheet name, grid of cell strings)] of every visible sheet of an Excel workbook.
    .xlsx/.xlsm need openpyxl, .xls needs xlrd. Formulas give their last saved value
    (a workbook generated by a program and never opened in Excel may have none)."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        return _read_xls(path)
    try:
        import openpyxl  # type: ignore
    except ImportError as e:
        raise RuntimeError("Para leer Excel instala openpyxl (pip install openpyxl).") from e
    import warnings
    full = path.stat().st_size <= _EXCEL_FULL_MODE_BYTES
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # "data validation extension not supported" ...
        wb = openpyxl.load_workbook(path, read_only=not full, data_only=True)
    out: list[tuple[str, list[list[str]]]] = []
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            grid = [[_excel_value(c.value, getattr(c, "number_format", "") or "") for c in row]
                    for row in ws.iter_rows()]
            if full:                              # merged cells: copy the top-left value
                for rng in ws.merged_cells.ranges:
                    r0, c0 = rng.min_row - 1, rng.min_col - 1
                    if r0 >= len(grid) or c0 >= len(grid[r0]):
                        continue
                    val = grid[r0][c0]
                    for r in range(r0, min(rng.max_row, len(grid))):
                        for c in range(c0, min(rng.max_col, len(grid[r]))):
                            if not grid[r][c]:
                                grid[r][c] = val
            while grid and not any(grid[-1]):
                grid.pop()
            out.append((ws.title, grid))
    finally:
        wb.close()
    return out


def _read_xls(path: Path) -> list[tuple[str, list[list[str]]]]:
    try:
        import xlrd  # type: ignore
    except ImportError as e:
        raise RuntimeError("Para leer Excel .xls (formato 97-2003) instala xlrd "
                           "(pip install xlrd) o guarda el archivo como .xlsx.") from e
    book = xlrd.open_workbook(str(path), formatting_info=False)
    out: list[tuple[str, list[list[str]]]] = []
    for sh in book.sheets():
        if getattr(sh, "visibility", 0) != 0:
            continue
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif cell.ctype in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                row.append(_excel_value(v))
            grid.append(row)
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []) or []:
            val = grid[r0][c0] if r0 < len(grid) and c0 < len(grid[r0]) else ""
            for r in range(r0, min(r1, len(grid))):
                for c in range(c0, min(c1, len(grid[r]))):
                    if not grid[r][c]:
                        grid[r][c] = val
        out.append((sh.name, grid))
    return out


def merge_chunks(texts: list[str]) -> str:
    """Join consecutive chunks of one file, removing the splitter's overlap."""
    out = ""
    for t in texts:
        t = t or ""
        if not out:
            out = t
            continue
        probe = t[:30]
        cut = -1
        if len(probe) >= 10:
            i = out.find(probe, max(0, len(out) - 450))
            while i != -1:
                if t.startswith(out[i:]):
                    cut = len(out) - i
                    break
                i = out.find(probe, i + 1)
        out = out + t[cut:] if cut > 0 else out + "\n" + t
    return out


@functools.lru_cache(maxsize=32)
def _load_image_png_cached(path: str, max_side: int, mtime: float) -> bytes:
    try:
        from PIL import Image  # type: ignore
    except ImportError as e:
        raise RuntimeError("Falta Pillow (pip install pillow).") from e
    Image.MAX_IMAGE_PIXELS = None
    with Image.open(path) as im:
        im.draft("RGB", (max_side, max_side))
        im = im.convert("RGB") if im.mode not in ("RGB", "L") else im.copy()
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, format="PNG", optimize=True)
        return buf.getvalue()


def load_image_png(path: str | Path, max_side: int | None = None) -> bytes:
    p = str(Path(path))
    return _load_image_png_cached(p, max_side or settings.image_max_side, os.path.getmtime(p))


def source_tag(d: dict[str, Any]) -> str:
    """'{MANUAL}', '{DATASHEET}' ... from the folder the chunk comes from (declared in the
    .env). Web pages and generated datasheets carry their own kind."""
    kind = d.get("kind")
    if kind == "web":
        return "{MANUFACTURER WEBSITE}"
    if kind == "datasheet":
        return "{DATASHEET (generated from product CSV)}"
    return "{" + SECTION_LABEL.get(d.get("section") or OTHER_SECTION, "OTHER DOCUMENT") + "}"


def format_docs(docs: list[dict[str, Any]], n: int = 6, width: int = 400) -> str:
    """Short listing for agent notes (drawings get more room for their parts list). Every
    line says which kind of source it is, so manuals, datasheets and drawings never mix
    silently."""
    if not docs:
        return "No matching documents in the local corpus."
    out = []
    for d in docs[:n]:
        kind = d.get("kind")
        w = 4000 if kind == "table" else (max(width, 900) if kind == "image" else width)
        out.append(f"[{d['id']}] {source_tag(d)} {d['text'][:w]}")
    return "\n".join(out)
